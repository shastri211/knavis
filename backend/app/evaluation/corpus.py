"""A small, deterministic evaluation corpus for a fictional company ("Northwind Retail").

Every document is generated from code, so ground truth is exact and nothing real is shared. The tabular files are
built from plain Python lists (``SALES_ROWS``, ``CAMPAIGN_ROWS``) that the dataset's expected numbers can be checked
against independently.
"""
import datetime as dt
import io

RETENTION_POLICY = """Northwind Retail - Data Retention Policy

1. Customer data must be retained for 90 days after the contract ends.
2. Backups are kept for 30 days and are then overwritten.
3. Audit logs are retained for 365 days in cold storage.
4. Deletion requests from customers are answered within 14 days.
5. Data under a legal hold is exempt from deletion until the legal team lifts the hold.
"""

OFFICE_NOTICE_HI = """कार्यालय सूचना

कार्यालय का समय सुबह 9:30 बजे से शाम 6:00 बजे तक है।
सभी कर्मचारियों को हर महीने की पहली तारीख तक अपनी उपस्थिति रिपोर्ट जमा करनी होगी।
दिवाली के लिए कार्यालय 3 दिन बंद रहेगा।
"""

MEETING_SRT = """1
00:00:02,000 --> 00:00:06,000
Welcome to the quarterly review.

2
00:00:07,000 --> 00:00:12,000
Revenue grew 18 percent compared to last quarter.

3
00:00:13,000 --> 00:00:18,000
The new warehouse opens in Pune in October.
"""

REGIONS = ["North", "South", "East", "West"]
PRODUCTS = {"Widget": 20, "Gadget": 35, "Gizmo": 50}          # unit prices
CHANNELS = ["Email", "Search", "Social", "Display"]

# (order id, date, region, product, units, revenue, returned 1/0)
SALES_ROWS = []
for _i in range(60):
    _product = list(PRODUCTS)[_i % 3]
    _units = 5 + (_i * 7) % 23
    SALES_ROWS.append((f"ORD-{1001 + _i}", dt.date(2025, 1, 1) + dt.timedelta(days=_i), REGIONS[_i % 4], _product,
                       _units, _units * PRODUCTS[_product], 1 if _i % 9 == 0 else 0))

# (campaign, channel, spend, clicks, converted 1/0)
CAMPAIGN_ROWS = []
for _i in range(40):
    CAMPAIGN_ROWS.append((f"CMP-{_i + 1:03d}", CHANNELS[(_i * 3) % 4], round(50 + (_i * 37) % 400 + 0.5, 2), 20 + (_i * 11) % 180,
                          1 if (_i * 5) % 7 < 3 else 0))


def _docx() -> bytes:
    from docx import Document
    doc = Document()
    doc.add_heading("Northwind Retail Employee Handbook", level=1)
    doc.add_heading("Leave", level=2)
    doc.add_paragraph("Full-time employees receive 24 days of paid annual leave per year.")
    doc.add_paragraph("Employees receive 10 days of paid sick leave per year. A doctor's note is required after 3 consecutive days of sick leave.")
    doc.add_heading("Expenses and equipment", level=2)
    doc.add_paragraph("Expense reports must be submitted within 30 days and are reimbursed within 10 business days.")
    doc.add_paragraph("Laptops are refreshed every 36 months.")
    doc.add_heading("Working arrangements", level=2)
    doc.add_paragraph("Employees may work remotely up to 3 days per week with manager approval.")
    doc.add_heading("Notice and probation", level=2)
    table = doc.add_table(rows=3, cols=3)
    for r, row in enumerate([("Level", "Notice period", "Probation"), ("Junior", "30 days", "3 months"), ("Senior", "60 days", "6 months")]):
        for c, cell in enumerate(row):
            table.cell(r, c).text = cell
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _pdf() -> bytes:
    import fitz
    pdf = fitz.open()
    page = pdf.new_page()
    text = ("Northwind Cloud Pricing 2025\n\n"
            "Basic plan: $12 per user per month.\n"
            "Pro plan: $29 per user per month.\n"
            "Enterprise plan: custom quote.\n\n"
            "Annual billing receives a 15% discount.\n"
            "The free trial lasts 14 days.\n"
            "Support is available Monday to Friday, 9am to 6pm IST.\n")
    page.insert_textbox(fitz.Rect(72, 72, 520, 400), text, fontsize=11)
    data = pdf.tobytes()
    pdf.close()
    return data


def _sales_xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q1"
    ws.append(["Order ID", "Date", "Region", "Product", "Units", "Revenue", "Returned"])
    for row in SALES_ROWS:
        ws.append(list(row))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _campaign_csv() -> bytes:
    lines = ["Campaign,Channel,Spend,Clicks,Converted"] + [",".join(str(v) for v in row) for row in CAMPAIGN_ROWS]
    return ("\n".join(lines) + "\n").encode()


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def build_corpus() -> dict[str, tuple[bytes, str]]:
    """File name -> (bytes, content type)."""
    return {
        "retention_policy.txt": (RETENTION_POLICY.encode(), "text/plain"),
        "employee_handbook.docx": (_docx(), DOCX),
        "pricing.pdf": (_pdf(), "application/pdf"),
        "office_notice_hi.txt": (OFFICE_NOTICE_HI.encode(), "text/plain"),
        "quarterly_review.srt": (MEETING_SRT.encode(), "application/x-subrip"),
        "sales_q1.xlsx": (_sales_xlsx(), XLSX),
        "campaign_results.csv": (_campaign_csv(), "text/csv"),
    }
