"""Find the figures worth describing (charts, graphs, diagrams, scans of handwriting) in a document.

Only candidates are collected here, locally: small images are dropped, duplicates are merged by
content hash, and the number kept is capped, largest first. Describing them is a paid call made
elsewhere, once per distinct image (cached by hash).
"""
import hashlib
import io
from dataclasses import dataclass
from pathlib import Path

import fitz

MIN_SIDE = 150            # px: smaller images are icons or bullets
MAX_SIDE = 1568           # px: larger images are scaled down (fewer tokens, same legibility)
PDF_IMAGE_MIN_SHARE = 0.05
PDF_DRAWING_MIN_SHARE = 0.08
PDF_DRAWING_MIN_PATHS = 8     # a page with fewer vector paths has no chart (rules and borders are thin: the area test drops them)


@dataclass
class Figure:
    key: str                  # element id relative to the document, e.g. "p3:fig0"
    locator: str
    png: bytes
    sha: str
    page: int | None = None
    slide: int | None = None
    area: float = 0.0


def normalize_png(data: bytes) -> tuple[bytes, int, int] | None:
    """PNG bytes at a sensible size, or None when the image is too small to be a figure."""
    from PIL import Image
    try:
        image = Image.open(io.BytesIO(data))
        image.load()
    except Exception:
        return None
    if max(image.size) < MIN_SIDE:
        return None
    if image.mode in ("RGBA", "LA", "P"):
        background = Image.new("RGB", image.size, "white")
        background.paste(image.convert("RGBA"), mask=image.convert("RGBA").split()[-1])
        image = background
    else:
        image = image.convert("RGB")
    if max(image.size) > MAX_SIDE:
        image.thumbnail((MAX_SIDE, MAX_SIDE))
    out = io.BytesIO()
    image.save(out, format="PNG")
    return out.getvalue(), image.width, image.height


def _make(key, locator, raw, page=None, slide=None) -> Figure | None:
    normalized = normalize_png(raw)
    if normalized is None:
        return None
    png, width, height = normalized
    return Figure(key=key, locator=locator, png=png, sha=hashlib.sha256(png).hexdigest(), page=page, slide=slide, area=width * height)


def _overlap(a: fitz.Rect, b: fitz.Rect) -> float:
    inter = a & b
    smaller = min(a.get_area(), b.get_area())
    return 0.0 if inter.is_empty or not smaller else inter.get_area() / smaller


def _pdf_figures(path: Path, info: dict, elements) -> list[Figure]:
    skip_pages = set(info.get("ocr_pages", []))   # scanned pages are read by OCR, not described
    table_boxes: dict[int, list] = {}
    for el in elements:
        if el.kind == "table" and el.bbox and el.page:
            table_boxes.setdefault(el.page, []).append(fitz.Rect(el.bbox))

    figures: list[Figure] = []
    with fitz.open(path) as doc:
        for number, page in enumerate(doc, 1):
            if number in skip_pages:
                continue
            page_area = page.rect.get_area() or 1.0
            candidates = []
            try:
                for image in page.get_image_info():
                    rect = fitz.Rect(image["bbox"])
                    if rect.get_area() / page_area >= PDF_IMAGE_MIN_SHARE:
                        candidates.append(rect)
            except Exception:
                pass
            try:
                if len(page.get_drawings()) >= PDF_DRAWING_MIN_PATHS:
                    for rect in page.cluster_drawings():
                        if rect.get_area() / page_area >= PDF_DRAWING_MIN_SHARE and not any(
                                _overlap(rect, t) > 0.5 for t in table_boxes.get(number, [])):
                            candidates.append(rect)   # a drawn chart; ruled tables were excluded above
            except Exception:
                pass
            accepted: list[fitz.Rect] = []
            for rect in sorted(candidates, key=lambda r: r.get_area(), reverse=True):
                if all(_overlap(rect, other) < 0.7 for other in accepted):
                    accepted.append(rect)
            for index, rect in enumerate(accepted):
                raw = page.get_pixmap(dpi=144, alpha=False, clip=rect).tobytes("png")
                figure = _make(f"p{number}:fig{index}", f"page {number}", raw, page=number)
                if figure:
                    figures.append(figure)
    return figures


def _docx_figures(path: Path) -> list[Figure]:
    from docx import Document
    document = Document(path)
    figures = []
    for rel in document.part.rels.values():
        if rel.reltype.endswith("/image"):
            figure = _make(f"fig{len(figures)}", f"figure {len(figures) + 1}", rel.target_part.blob)
            if figure:
                figures.append(figure)
    return figures


def _pptx_figures(path: Path) -> list[Figure]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def pictures(shapes):
        for shape in shapes:
            try:
                kind = shape.shape_type
            except NotImplementedError:
                continue
            if kind == MSO_SHAPE_TYPE.PICTURE:
                yield shape
            elif kind == MSO_SHAPE_TYPE.GROUP:
                yield from pictures(shape.shapes)

    figures = []
    for number, slide in enumerate(Presentation(path).slides, 1):
        for index, shape in enumerate(pictures(slide.shapes)):
            figure = _make(f"s{number}:fig{index}", f"slide {number}", shape.image.blob, slide=number)
            if figure:
                figures.append(figure)
    return figures


def collect_figures(path: Path, kind: str, info: dict, elements, limit: int) -> list[Figure]:
    """Candidate figures of a document: distinct, at least MIN_SIDE px, largest first, at most ``limit``."""
    path = Path(path)
    if kind == "pdf":
        found = _pdf_figures(path, info, elements)
    elif kind == "docx":
        found = _docx_figures(path)
    elif kind == "pptx":
        found = _pptx_figures(path)
    else:
        return []
    distinct = {f.sha: f for f in found}.values()          # the same logo on every page is one figure
    return sorted(distinct, key=lambda f: f.area, reverse=True)[:limit]
