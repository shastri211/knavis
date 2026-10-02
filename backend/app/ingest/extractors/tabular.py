"""Spreadsheets and delimited files.

A sheet becomes one summary element (size, columns, types, an example row) plus row groups
that each repeat the header, so any retrieved group is self-describing. Nothing is sent
to a model.
"""
import csv
import datetime as dt
import re
from dataclasses import dataclass
from pathlib import Path

from ..elements import PARAGRAPH, SUMMARY, TABLE, Element, ExtractionError, ExtractionResult
from ..tables import clean_rows_numbered, render_rows, split_rows

MAX_ROWS = 100_000          # rows read per sheet; a longer sheet is truncated and flagged
GROUP_CHARS = 1200          # size of one row group (about one chunk)
_TYPE_SAMPLE = 200
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}([ T]\d{2}:\d{2}(:\d{2})?)?$")


def decode_bytes(raw: bytes) -> str:
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def _cell(value):
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, dt.date):
        return value.isoformat()
    return value


def _column_type(values: list[str]) -> str:
    values = [v for v in values if v]
    if not values:
        return "empty"
    if all(_DATE_RE.match(v) for v in values):
        return "date"
    try:
        for v in values:
            float(v.replace(",", ""))
        return "number"
    except ValueError:
        return "text"


def _is_number(value: str) -> bool:
    try:
        float(value.replace(",", ""))
        return True
    except ValueError:
        return False


def _header_index(rows: list[list[str]]) -> int:
    """Index of the header row. Sheets often open with a title or notes: the header is the
    first of the opening rows that is nearly as wide as the widest one and mostly text."""
    probe = rows[:10]
    counts = [sum(1 for c in r if c) for r in probe]
    widest = max(counts)
    if widest < 2:
        return 0
    for index, row in enumerate(probe):
        filled = counts[index]
        if filled >= max(2, 0.6 * widest) and sum(1 for c in row if c and not _is_number(c)) >= 0.6 * filled:
            return index
    return 0


def sheet_elements(rows, *, sheet: str | None, prefix: str, truncated: bool = False) -> tuple[list[Element], dict]:
    rows, numbers = clean_rows_numbered(rows)   # numbers = the rows' real positions in the sheet
    if not rows:
        return [], {"rows": 0, "columns": 0}
    where = f"Sheet '{sheet}'" if sheet else "Table"
    place = f"sheet {sheet}, " if sheet else ""

    elements: list[Element] = []
    start = _header_index(rows)
    if start:   # title or notes above the header
        preamble = "\n".join(" | ".join(c for c in row if c) for row in rows[:start])
        elements.append(Element(
            id=f"{prefix}:pre", kind=PARAGRAPH, text=preamble, sheet=sheet,
            locator=f"{place}rows {numbers[0]}-{numbers[start - 1]}" if start > 1 else f"{place}row {numbers[0]}",
        ))
    rows, numbers = rows[start:], numbers[start:]

    header, data = rows[0], rows[1:]
    names = [h or f"column {i}" for i, h in enumerate(header, 1)]
    rows[0] = names
    types = [_column_type([r[i] for r in data[:_TYPE_SAMPLE]]) for i in range(len(names))]
    summary = (
        f"{where}: {len(data)} data rows, {len(names)} columns"
        f"{' (truncated: only the first %d rows were read)' % MAX_ROWS if truncated else ''}.\n"
        f"Columns: {', '.join(f'{n} ({t})' for n, t in zip(names, types))}."
    )
    if data:
        summary += f"\nExample row: {' | '.join(data[0])}"
    elements.append(Element(
        id=f"{prefix}:summary", kind=SUMMARY, text=summary, sheet=sheet,
        locator=f"sheet {sheet}" if sheet else "table",
        meta={"rows": len(data), "columns": names, "types": types, "truncated": truncated, "header_row": numbers[0]},
    ))
    for first, last, group in split_rows(rows, GROUP_CHARS):
        a, b = numbers[first - 1], numbers[last - 1]   # real sheet row numbers; blank rows were skipped
        elements.append(Element(
            id=f"{prefix}:r{a}", kind=TABLE, text=render_rows(group), rows=group, sheet=sheet,
            locator=f"{place}rows {a}-{b}" if b > a else f"{place}row {a}",
            meta={"row_start": a, "row_end": b},
        ))
    return elements, {"rows": len(data), "columns": len(names), "truncated": truncated}


def _read_xlsx(path: Path) -> list[tuple[str, list[list], bool]]:
    """``(sheet name, raw rows, truncated)`` for every sheet of a workbook."""
    import openpyxl
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError("This spreadsheet could not be opened; it may be corrupted or password protected.") from exc
    sheets = []
    try:
        for ws in workbook.worksheets:
            rows = []
            for count, row in enumerate(ws.iter_rows(values_only=True)):
                if count > MAX_ROWS + 1:
                    break
                rows.append([_cell(v) for v in row])
            sheets.append((ws.title, rows[:MAX_ROWS + 1], len(rows) > MAX_ROWS + 1))
    finally:
        workbook.close()
    return sheets


def _read_csv(path: Path) -> tuple[list[list], bool]:
    text = decode_bytes(path.read_bytes())
    if path.suffix.lower() == ".tsv":
        dialect = csv.excel_tab
    else:
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
    csv.field_size_limit(10_000_000)
    rows = []
    for count, row in enumerate(csv.reader(text.splitlines(), dialect)):
        if count > MAX_ROWS + 1:
            break
        rows.append(row)
    return rows[:MAX_ROWS + 1], len(rows) > MAX_ROWS + 1


def extract_xlsx(path: Path) -> ExtractionResult:
    elements, sheets, tables = [], [], []
    for index, (title, rows, truncated) in enumerate(_read_xlsx(path)):
        sheet_els, stats = sheet_elements(rows, sheet=title, prefix=f"sh{index}", truncated=truncated)
        elements.extend(sheet_els)
        sheets.append({"name": title, **stats})
        if table := _sheet_table(rows, title, truncated):
            tables.append(table)
    return ExtractionResult("xlsx", elements, {"sheets": sheets}, tables)


def extract_csv(path: Path) -> ExtractionResult:
    rows, truncated = _read_csv(path)
    elements, stats = sheet_elements(rows, sheet=None, prefix="t", truncated=truncated)
    table = _sheet_table(rows, None, truncated)
    return ExtractionResult("csv", elements, {"sheets": [{"name": path.stem, **stats}]}, [table] if table else [])


# ---- whole sheets as data (the analytics table store) -----------------------------------------

@dataclass
class SheetTable:
    """One sheet's data rows exactly as the extractor sees them, with each row's real sheet number."""
    sheet: str | None
    header: list[str]
    rows: list[list[str]]       # data rows, as text, as wide as the header
    numbers: list[int]          # the sheet row number of each data row
    truncated: bool = False


def _sheet_table(rows, sheet: str | None, truncated: bool) -> SheetTable | None:
    rows, numbers = clean_rows_numbered(rows)
    if not rows:
        return None
    start = _header_index(rows)
    rows, numbers = rows[start:], numbers[start:]
    header = [h or f"column {i}" for i, h in enumerate(rows[0], 1)]
    if len(rows) < 2:
        return None   # a header with no data is not a table
    return SheetTable(sheet, header, rows[1:], numbers[1:], truncated)


def read_tables(path: Path) -> list[SheetTable]:
    """Every sheet of an .xlsx/.xlsm workbook or a delimited text file as a table. Local and free."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        tables = [_sheet_table(rows, title, truncated) for title, rows, truncated in _read_xlsx(path)]
    else:
        rows, truncated = _read_csv(path)
        tables = [_sheet_table(rows, None, truncated)]
    return [t for t in tables if t]
