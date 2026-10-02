"""DOCX and PPTX: document order, headings, tables, notes, and image counts for later vision."""
import re
from pathlib import Path

from ..elements import FIGURE, HEADING, PARAGRAPH, SLIDE, TABLE, Element, ExtractionError, ExtractionResult
from ..tables import clean_cell, clean_rows, render_rows

_HEADING_STYLE_RE = re.compile(r"^heading\s*(\d+)$", re.IGNORECASE)


def _docx_heading_level(style_name: str) -> int | None:
    name = (style_name or "").strip()
    if name.lower() == "title":
        return 1
    match = _HEADING_STYLE_RE.match(name)
    return min(int(match.group(1)), 6) if match else None


def extract_docx(path: Path) -> ExtractionResult:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(path)
    except Exception as exc:
        raise ExtractionError("This Word document could not be opened; it may be corrupted.") from exc

    elements: list[Element] = []
    figures = tables = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        position = len(elements)
        if tag == "p":
            paragraph = Paragraph(child, document)
            figures += len(child.findall(".//{*}drawing"))
            text = re.sub(r"\s+", " ", paragraph.text).strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style is not None else ""
            level = _docx_heading_level(style)
            if level:
                elements.append(Element(id=f"e{position}", kind=HEADING, text=text, level=level))
                continue
            if style.lower().startswith("list"):
                text = f"- {text}"
            elements.append(Element(id=f"e{position}", kind=PARAGRAPH, text=text))
        elif tag == "tbl":
            table = Table(child, document)
            rows = []
            for row in table.rows:
                seen, cells = set(), []
                for cell in row.cells:
                    if cell._tc in seen:   # merged cells repeat; keep one copy
                        continue
                    seen.add(cell._tc)
                    cells.append(cell.text)
                rows.append(cells)
            rows = clean_rows(rows)
            if rows:
                tables += 1
                elements.append(Element(id=f"e{position}", kind=TABLE, text=render_rows(rows), rows=rows, locator=f"table {tables}"))

    info = {"tables": tables, "figures": figures, "estimated_vision_calls": figures}
    return ExtractionResult("docx", elements, info)


def _shape_type(shape):
    try:
        return shape.shape_type
    except NotImplementedError:   # some placeholder shapes have no type
        return None


def _shapes(shapes):
    """All shapes of a slide, descending into groups."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for shape in shapes:
        yield shape
        if _shape_type(shape) == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes)


def _chart_text(chart) -> str:
    """Title, type, and the plotted series as a table, straight from the chart object (no model call)."""
    try:
        title = chart.chart_title.text_frame.text.strip() if chart.has_title else ""
    except Exception:
        title = ""
    try:
        kind = chart.chart_type.name.replace("_", " ").lower()
    except Exception:
        kind = "chart"
    rows: list[list[str]] = []
    try:
        for plot in chart.plots:
            categories = [str(c) for c in plot.categories]
            series = [(s.name or f"series {i}", list(s.values)) for i, s in enumerate(plot.series, 1)]
            if not rows:
                rows.append(["category"] + [name for name, _ in series])
            for index, category in enumerate(categories):
                rows.append([category] + [("" if index >= len(v) or v[index] is None else clean_cell(v[index])) for _, v in series])
    except Exception:
        pass
    rows = clean_rows(rows)
    head = f"Chart ({kind}){': ' + title if title else ''}"
    return f"{head}\n{render_rows(rows)}" if len(rows) > 1 else ""


def extract_pptx(path: Path) -> ExtractionResult:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        presentation = Presentation(path)
    except Exception as exc:
        raise ExtractionError("This presentation could not be opened; it may be corrupted.") from exc

    elements: list[Element] = []
    figure_slides, picture_only_slides, tables, charts = [], [], 0, 0
    for number, slide in enumerate(presentation.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        parts, pictures = [], 0
        for shape in _shapes(slide.shapes):
            if _shape_type(shape) == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_chart", False) and shape.has_chart:
                # A native chart carries its data: read it locally instead of describing a picture of it.
                chart_text = _chart_text(shape.chart)
                if chart_text:
                    charts += 1
                    elements.append(Element(
                        id=f"s{number}:chart{charts}", kind=FIGURE, text=chart_text, slide=number,
                        locator=f"slide {number}, chart", source="native",
                    ))
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = clean_rows([[cell.text for cell in row.cells] for row in shape.table.rows])
                if rows:
                    tables += 1
                    elements.append(Element(
                        id=f"s{number}:t{tables}", kind=TABLE, text=render_rows(rows), rows=rows,
                        slide=number, locator=f"slide {number}",
                    ))
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
                if text and text != title:
                    parts.append(text)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        body = "\n".join(parts)
        pieces = [title, body] + ([f"Speaker notes: {notes}"] if notes else [])
        text = "\n".join(p for p in pieces if p)
        if text:
            elements.append(Element(
                id=f"s{number}", kind=SLIDE, text=text, slide=number, locator=f"slide {number}",
                meta={"title": title} if title else {},
            ))
        if pictures:
            figure_slides.append({"slide": number, "count": pictures})
            if not text:
                picture_only_slides.append(number)

    info = {
        "slides": len(presentation.slides), "tables": tables, "figure_slides": figure_slides,
        "picture_only_slides": picture_only_slides,
        "estimated_vision_calls": sum(f["count"] for f in figure_slides),
    }
    return ExtractionResult("pptx", elements, info)
