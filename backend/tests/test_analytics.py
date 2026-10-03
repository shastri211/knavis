"""Spreadsheet analytics: the table store, the SQL guard, deterministic detection, and the golden end-to-end path
(real app, real xlsx/csv ingestion, real SQLite execution; only the model that writes the SQL is faked)."""
from types import SimpleNamespace

import pytest

from conftest import CAMPAIGN_FILE, CAMPAIGN_TABLE, XLSX_TYPE, make_campaign_xlsx, make_txt

WHICH_CHANNEL = "Which channel had the highest number of conversions?"
BEST_SQL = (f'SELECT "Channel", SUM("Converted") AS conversions FROM "{CAMPAIGN_TABLE}" '
            'GROUP BY "Channel" ORDER BY conversions DESC LIMIT 5')


# ---- typing and naming ---------------------------------------------------------------------

def test_number_parsing_keeps_codes_and_ids_as_text():
    from app.analytics.tablestore import parse_number
    assert parse_number("1,234") == 1234 and parse_number("$5.50") == 5.5 and parse_number("-3") == -3
    assert parse_number("007") is None and parse_number("12abc") is None and parse_number("1.2.3") is None


def test_column_inference():
    from app.analytics.tablestore import infer_column
    assert infer_column(["1", "0", "", "1"]) == ("integer", [1, 0, None, 1])
    assert infer_column(["1.5", "2"])[0] == "number"
    assert infer_column(["2024-01-05", "2024-02-01"])[0] == "date"
    assert infer_column(["TRUE", "false"]) == ("boolean", [1, 0])
    assert infer_column(["N/A", "-", ""])[0] == "empty"
    assert infer_column(["10", "ten"])[0] == "text"           # one bad cell must not silently drop data from a sum
    assert infer_column(["001", "002"])[0] == "text"


def test_column_names_are_unique_safe_and_keep_hindi_words_whole():
    from app.analytics.tablestore import column_names
    assert column_names(["Conversion Rate", "conversion rate", "", "2024 Sales", "Row"]) == [
        "Conversion_Rate", "conversion_rate_2", "column_3", "c_2024_Sales", "Row"]
    assert column_names(["बिक्री (₹)"]) == ["बिक्री"]


# ---- SQL guard -----------------------------------------------------------------------------

@pytest.fixture
def table_session(client, session_id, upload):
    document, _ = upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    assert document["status"] == "indexed" and document["details"]["tables"][0]["name"] == CAMPAIGN_TABLE
    from app.analytics.tablestore import session_tables
    from app.db import SessionLocal
    with SessionLocal() as db:
        tables = session_tables(db, session_id)
        for t in tables:
            db.expunge(t)
    return session_id, tables


def _run(table_session, sql, **kw):
    from app.analytics.sqlguard import run_select
    session_id, tables = table_session
    return run_select(session_id, sql, tables, **kw)


def test_guard_runs_a_select_and_reports_exactly_which_columns_it_read(table_session):
    result = _run(table_session, BEST_SQL + ";")
    assert result.rows[0] == ("Email", 7) and result.rows[1] == ("Search", 5)
    assert result.columns_used == {CAMPAIGN_TABLE: {"Channel", "Converted"}}


def test_guard_supports_cte_subquery_and_count_star(table_session):
    assert _run(table_session, f'SELECT COUNT(*) FROM "{CAMPAIGN_TABLE}"').rows == [(30,)]
    assert _run(table_session, f'WITH t AS (SELECT * FROM "{CAMPAIGN_TABLE}") SELECT MAX("_row") FROM t').rows == [(31,)]
    assert _run(table_session, f'SELECT (SELECT MIN("Cost") FROM "{CAMPAIGN_TABLE}")').rows[0][0] == 11.5


@pytest.mark.parametrize("sql", [
    f'DROP TABLE "{CAMPAIGN_TABLE}"',
    f'DELETE FROM "{CAMPAIGN_TABLE}"',
    f'UPDATE "{CAMPAIGN_TABLE}" SET "Converted" = 1',
    f'INSERT INTO "{CAMPAIGN_TABLE}" VALUES (1, 1, 1, 1, 1, 1)',
    f'CREATE TABLE x AS SELECT * FROM "{CAMPAIGN_TABLE}"',
    "PRAGMA table_info(x)",
    "ATTACH DATABASE 'other.db' AS o",
    f'SELECT 1; SELECT * FROM "{CAMPAIGN_TABLE}"',
    f'SELECT * FROM "{CAMPAIGN_TABLE}" -- hidden',
    f'SELECT /* hi */ * FROM "{CAMPAIGN_TABLE}"',
    "SELECT * FROM sqlite_master",
    "SELECT * FROM sessions",
    f'SELECT "password" FROM "{CAMPAIGN_TABLE}"',
    f'SELECT load_extension("x") FROM "{CAMPAIGN_TABLE}"',
    f'SELECT randomblob(1000000000) FROM "{CAMPAIGN_TABLE}"',
    "WITH RECURSIVE r(n) AS (SELECT 1 UNION ALL SELECT n + 1 FROM r) SELECT n FROM r",
    "EXPLAIN SELECT 1",
    "",
    "   ",
    None,
    "SELECT " + "1," * 2000 + "1",
])
def test_guard_refuses_everything_that_is_not_a_plain_allowed_select(table_session, sql):
    from app.analytics.sqlguard import SqlRejected
    with pytest.raises(SqlRejected):
        _run(table_session, sql)


def test_guard_keeps_the_data_intact_after_attempted_writes(table_session):
    from app.analytics.sqlguard import SqlRejected
    for sql in (f'DROP TABLE "{CAMPAIGN_TABLE}"', f'DELETE FROM "{CAMPAIGN_TABLE}"'):
        with pytest.raises(SqlRejected):
            _run(table_session, sql)
    assert _run(table_session, f'SELECT COUNT(*) FROM "{CAMPAIGN_TABLE}"').rows == [(30,)]


def test_guard_stops_a_runaway_query_and_caps_rows(table_session):
    from app.analytics.sqlguard import SqlRejected
    t = f'"{CAMPAIGN_TABLE}"'
    with pytest.raises(SqlRejected, match="too long"):
        _run(table_session, f"SELECT COUNT(*) FROM {t} a, {t} b, {t} c, {t} d, {t} e, {t} f, {t} g", timeout=0.3)
    capped = _run(table_session, f'SELECT * FROM {t}', max_rows=5)
    assert len(capped.rows) == 5 and capped.truncated


def test_guard_reads_quotes_properly(table_session):
    """A ';' or '--' inside a string literal is data; a doubled quote does not end the literal."""
    t = f'"{CAMPAIGN_TABLE}"'
    assert _run(table_session, f"SELECT COUNT(*) FROM {t} WHERE \"Channel\" = 'a;b -- c'").rows == [(0,)]
    assert _run(table_session, f"SELECT COUNT(*) FROM {t} WHERE \"Channel\" = 'it''s'").rows == [(0,)]


def test_a_misspelt_column_is_an_error_not_a_string_literal(table_session):
    """SQLite's double-quoted-string fallback would otherwise return the wrong name as if it were data."""
    from app.analytics.sqlguard import SqlRejected
    with pytest.raises(SqlRejected, match="not a table or column"):
        _run(table_session, f'SELECT "Revnue", SUM("Converted") FROM "{CAMPAIGN_TABLE}" GROUP BY 1')


def test_one_session_cannot_read_anothers_tables(client, upload, table_session):
    from app.analytics.sqlguard import SqlRejected, run_select
    other = client.post("/api/sessions", json={"title": "other"}).json()["id"]
    upload(other, "other.csv", b"a,b\n1,2\n3,4\n", "text/csv")
    _, first_tables = table_session
    with pytest.raises(SqlRejected):
        run_select(other, f'SELECT * FROM "{CAMPAIGN_TABLE}"', [SimpleNamespace(table_name="other", columns_json=[{"name": "a"}, {"name": "b"}])])


# ---- detection -----------------------------------------------------------------------------

@pytest.fixture
def campaign_tables():
    columns = [
        {"name": "Campaign_ID", "original": "Campaign_ID", "kind": "text", "type": "TEXT"},
        {"name": "Channel", "original": "Channel", "kind": "text", "type": "TEXT", "values": ["Email", "Search", "Social"]},
        {"name": "Cost", "original": "Cost", "kind": "number", "type": "REAL"},
        {"name": "Converted", "original": "Converted", "kind": "integer", "type": "INTEGER", "values": ["0", "1"]},
    ]
    return [SimpleNamespace(table_name=CAMPAIGN_TABLE, sheet="Week1", filename=CAMPAIGN_FILE, columns_json=columns)]


@pytest.mark.parametrize("question", [
    "Which channel had the highest number of conversions?",
    "total cost by channel",
    "How many rows are there?",
    "what is the average cost",
    "How many campaigns converted from Email?",
    "list rows where channel is Search",
    "सबसे ज्यादा conversions किस channel में हुए",
    "kul cost kitna hai",
    "Which channel had the lowest cost?",
])
def test_analytical_questions_are_detected(question, campaign_tables):
    from app.analytics.detect import analyze
    assert analyze(question, campaign_tables, has_other_documents=True) is not None


@pytest.mark.parametrize("question", [
    "How long must company data be retained after the contract ends?",
    "How many days are backups kept?",
    "Summarize the report",
    "What is the refund policy?",
    "Who is the author?",
])
def test_ordinary_document_questions_are_left_to_retrieval(question, campaign_tables):
    from app.analytics.detect import analyze
    assert analyze(question, campaign_tables, has_other_documents=True) is None


def test_ranking_direction_and_named_columns(campaign_tables):
    from app.analytics.detect import analyze
    top = analyze(WHICH_CHANNEL, campaign_tables, True)
    assert top.ranking == "highest" and (CAMPAIGN_TABLE, "Channel") in top.columns and (CAMPAIGN_TABLE, "Converted") in top.columns
    assert analyze("which channel had the lowest cost", campaign_tables, True).ranking == "lowest"
    assert analyze("campaigns with at least 5 conversions by channel", campaign_tables, True).ranking is None


def test_with_only_tables_a_computation_needs_no_column_match(campaign_tables):
    from app.analytics.detect import analyze
    assert analyze("what is the grand total?", campaign_tables, has_other_documents=False) is not None
    assert analyze("what is the grand total?", campaign_tables, has_other_documents=True) is None


# ---- rendering and verification ------------------------------------------------------------

def test_label_numbers_must_come_from_the_result_or_the_question():
    from app.analytics.engine import verified_label
    rows = [("Email", 7), ("Search", 5)]
    assert verified_label("Conversions by channel", "q", ["c", "n"], rows) == "Conversions by channel"
    assert verified_label("Top 2 channels", "q", ["c", "n"], rows) == "Top 2 channels"          # 2 = number of rows
    assert verified_label("Email converted 999 times", "q", ["c", "n"], rows) == ""
    assert verified_label("Sales above 100", "sales above 100?", ["c", "n"], rows) == "Sales above 100"


def test_number_formatting_is_exact():
    from app.analytics.engine import format_value
    assert format_value(7) == "7" and format_value(7.0) == "7" and format_value(0.333333) == "0.3333"
    assert format_value(1234567) == "1,234,567" and format_value(None) == "" and format_value("x") == "x"


# ---- golden end to end ---------------------------------------------------------------------

def test_spreadsheet_is_loaded_as_a_queryable_table_with_a_catalog(client, upload, session_id):
    document, job = upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(notes=True), XLSX_TYPE)
    assert job["status"] == "completed"
    [table] = document["details"]["tables"]
    assert table == {"name": CAMPAIGN_TABLE, "sheet": "Week1", "rows": 30, "columns": 5}
    from conftest import table_files
    assert len(table_files(session_id)) == 1


def test_which_channel_had_the_highest_conversions_is_computed_not_guessed(llm, session_id, upload, ask):
    upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    llm.sql, llm.sql_label = BEST_SQL, "Conversions by channel, highest first"
    result = ask(session_id, WHICH_CHANNEL)

    answer = result["message"]["content"]
    assert answer.startswith("Email has the highest conversions: 7.")
    assert "Search | 5" in answer and "Social | 3" in answer
    assert llm.calls == ["sql"]                                   # one model call: the SQL writer; the answer is rendered from the rows
    assert "Converted" in llm.sql_prompt and "flag: 1 = yes" in llm.sql_prompt and f'"{CAMPAIGN_TABLE}"' in llm.sql_prompt
    [citation] = result["citations"]
    assert citation["source"] == CAMPAIGN_FILE and citation["sheet"] == "Week1"
    assert citation["columns"] == ["Channel", "Converted"] and "all 30 data rows (2-31)" in citation["locator"]
    assert f"Source: {CAMPAIGN_FILE}, sheet Week1, all 30 data rows (2-31)." in answer
    assert result["route"] == "rag"


def test_the_executed_sql_is_recorded_and_usage_is_billed(llm, session_id, upload, ask):
    from app.db import SessionLocal
    from app.models import UsageEvent
    upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    llm.sql = BEST_SQL
    ask(session_id, WHICH_CHANNEL)
    with SessionLocal() as db:
        assert db.query(UsageEvent).filter(UsageEvent.session_id == session_id).count() == 1


def test_scalar_questions_get_a_scalar_answer(llm, session_id, upload, ask):
    upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    llm.sql, llm.sql_label = f'SELECT COUNT(*) AS row_count FROM "{CAMPAIGN_TABLE}"', "Number of rows"
    answer = ask(session_id, "How many rows are in the sheet?")["message"]["content"]
    assert answer.startswith("Number of rows: 30")


def test_a_filtered_row_listing_cites_the_exact_sheet_rows(llm, session_id, upload, ask):
    upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    llm.sql = f'SELECT "_row", "Campaign_ID", "Channel" FROM "{CAMPAIGN_TABLE}" WHERE "Channel" = \'Social\' LIMIT 3'
    result = ask(session_id, "list rows where channel is Social")
    assert "rows 22-24" in result["citations"][0]["locator"]   # sheet rows, header is row 1, Social starts after 20 data rows
    assert "C021" in result["message"]["content"]


def test_a_label_with_an_invented_number_is_dropped_but_the_computed_answer_stays(llm, session_id, upload, ask):
    upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    llm.sql, llm.sql_label = f'SELECT SUM("Converted") AS total_conversions FROM "{CAMPAIGN_TABLE}"', "Email alone converted 4000 times"
    answer = ask(session_id, "What is the total number of conversions?")["message"]["content"]
    assert "4000" not in answer and "Total conversions: 15" in answer


@pytest.mark.parametrize("sql", [
    f'DROP TABLE "{CAMPAIGN_TABLE}"',
    f'SELECT "Revenue" FROM "{CAMPAIGN_TABLE}"',
    "SELECT * FROM missing_table",
    f'SELECT * FROM "{CAMPAIGN_TABLE}"; DROP TABLE "{CAMPAIGN_TABLE}"',
    "this is not sql",
])
def test_an_invalid_query_abstains_and_never_guesses(sql, llm, session_id, upload, ask):
    upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    llm.sql = sql
    result = ask(session_id, WHICH_CHANNEL)
    assert "couldn't compute that reliably" in result["message"]["content"] and result["citations"] == []
    assert llm.calls == ["sql"]
    llm.sql = BEST_SQL                                            # the data survived
    assert ask(session_id, WHICH_CHANNEL)["message"]["content"].startswith("Email has the highest")


def test_a_malformed_model_reply_abstains(llm, session_id, upload, ask):
    upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    for raw in ("I think it is Email", "{not json", ""):
        llm.sql_raw = raw
        assert "couldn't compute" in ask(session_id, WHICH_CHANNEL)["message"]["content"]


def test_a_reply_wrapped_in_a_code_fence_is_accepted(llm, session_id, upload, ask):
    upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    llm.sql_raw = '```json\n{"sql": "' + BEST_SQL.replace('"', '\\"') + '", "label": "x"}\n```'
    assert ask(session_id, WHICH_CHANNEL)["message"]["content"].startswith("Email has the highest")


def test_an_ordinary_question_in_a_spreadsheet_session_does_not_call_the_sql_writer(llm, session_id, upload, ask):
    upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    upload(session_id, "policy.txt", make_txt(), "text/plain")
    ask(session_id, "How long must company data be retained after the contract ends?")
    assert llm.calls == ["answer"]


def test_when_the_writer_says_it_is_not_a_table_question_retrieval_gets_a_turn(llm, session_id, upload, ask):
    """The wording looks analytical ("how many channels") but the text of a document answers it."""
    upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    upload(session_id, "policy.txt", make_txt(), "text/plain")
    llm.sql = None
    result = ask(session_id, "How many channels does the retention policy mention for data after the contract ends?")
    assert llm.calls == ["sql", "answer"]
    assert result["message"]["content"] == llm.answer and result["citations"][0]["source"] == "policy.txt"


def test_a_hindi_question_over_a_spreadsheet_only_session_is_computed(llm, session_id, upload, ask):
    upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    llm.sql, llm.sql_label = BEST_SQL, "चैनल के अनुसार कन्वर्शन"
    answer = ask(session_id, "सबसे ज्यादा conversions किस channel में हुए?")["message"]["content"]
    assert "Email" in answer and llm.calls == ["sql"]


def test_csv_files_are_queryable_too(llm, session_id, upload, ask):
    upload(session_id, "orders.csv", b"region,sales\nNorth,100\nSouth,250\nNorth,50\n", "text/csv")
    llm.sql, llm.sql_label = 'SELECT "region", SUM("sales") AS total_sales FROM "orders" GROUP BY "region" ORDER BY total_sales DESC', "Sales by region"
    answer = ask(session_id, "total sales by region")["message"]["content"]
    assert "South | 250" in answer and "North | 150" in answer and "orders.csv" in answer


def test_reuploading_replaces_tables_instead_of_duplicating_them(client, session_id, upload):
    from app.db import SessionLocal
    from app.models import DataTable
    document, _ = upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    assert client.post(f"/api/documents/{document['id']}/process", json={"action": "reindex"}).status_code == 200
    with SessionLocal() as db:
        assert db.query(DataTable).filter(DataTable.session_id == session_id).count() == 1


def test_the_sheet_summary_is_still_searchable_as_text(llm, session_id, upload, ask):
    """Lookups keep working through retrieval; only computation moved to SQL."""
    upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    llm.answer = "The sheet has 30 data rows [EVIDENCE 1]."
    result = ask(session_id, "What columns does the Week1 sheet have?")
    assert llm.calls == ["answer"] and result["citations"]


def test_deleting_a_document_drops_its_tables_index_rows_and_file(client, session_id, upload):
    from conftest import table_files
    from app.db import SessionLocal
    from app.models import DataTable, DocChunk
    from app.retrieval import fts
    document, _ = upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    assert client.delete(f"/api/documents/{document['id']}").status_code == 204
    with SessionLocal() as db:
        assert db.query(DataTable).filter(DataTable.session_id == session_id).count() == 0
        assert db.query(DocChunk).filter(DocChunk.document_id == document["id"]).count() == 0
        assert fts.search(db, session_id, "campaign channel", 5) == []
    assert not table_files(session_id)
    assert client.delete(f"/api/documents/{document['id']}").status_code == 404


# ---- answer wording ------------------------------------------------------------------------

def _lead(ranking, columns, rows, label=""):
    from app.analytics.detect import Intent
    from app.analytics.engine import _lead_sentence
    return _lead_sentence(Intent(strong=True, ranking=ranking), label, columns, rows)


def test_the_lead_sentence_names_the_winner_only_when_the_rows_are_really_ordered():
    cols = ["channel", "total_conversions"]
    assert _lead("highest", cols, [("Email", 7), ("Search", 5)]) == "Email has the highest total conversions: 7."
    assert _lead("lowest", cols, [("Social", 3), ("Search", 5)]) == "Social has the lowest total conversions: 3."
    assert _lead("highest", cols, [("Email", 7), ("Search", 7), ("Social", 3)]) == "Email and Search are tied for the highest total conversions: 7."
    # the model ordered ascending for a "highest" question: no claim is made, the table speaks for itself
    assert _lead("highest", cols, [("Social", 3), ("Email", 7)], label="Conversions") == "Conversions"


def test_empty_and_null_results_say_so_instead_of_inventing_a_number():
    assert _lead(None, ["total"], []) == "No matching rows were found."
    assert _lead(None, ["total"], [(None,)]) == "No matching rows were found."
    assert _lead(None, ["total"], [(0,)]) == "Total: 0"


def test_long_results_are_shortened_and_say_so():
    from app.analytics.engine import _table_text
    text = _table_text(["a"], [(i,) for i in range(50)], 50, truncated=False)
    assert text.splitlines()[-1] == "(showing the first 20 of 50 result rows)" and len(text.splitlines()) == 22


def test_row_numbers_are_compressed_into_ranges():
    from app.analytics.engine import _compress
    assert _compress([5, 3, 4, 9, 11, 10, 3]) == "3-5, 9-11"


# ---- switching it off ----------------------------------------------------------------------

def test_analytics_can_be_turned_off(llm, session_id, upload, ask, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "analytics_enabled", False)
    upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(), XLSX_TYPE)
    llm.sql = BEST_SQL
    ask(session_id, WHICH_CHANNEL)
    assert "sql" not in llm.calls                                  # retrieval answered (or abstained) as before


def test_the_result_row_limit_is_a_setting(table_session, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "analytics_max_rows", 4)
    result = _run(table_session, f'SELECT * FROM "{CAMPAIGN_TABLE}"')
    assert len(result.rows) == 4 and result.truncated


def test_unknown_quoted_names_are_refused_on_every_python_version(table_session, monkeypatch):
    """Without sqlite3's setconfig (Python < 3.12) the static check still stops "Revnue" from turning into a string."""
    import sqlite3
    from app.analytics.sqlguard import SqlRejected

    class NoSetconfig:
        def __init__(self, connection): self._c = connection
        def __getattr__(self, name):
            if name == "setconfig":
                raise AttributeError(name)
            return getattr(self._c, name)

    real = sqlite3.connect
    monkeypatch.setattr(sqlite3, "connect", lambda *a, **k: NoSetconfig(real(*a, **k)))
    with pytest.raises(SqlRejected, match="not a table or column"):
        _run(table_session, f'SELECT "Revnue", SUM("Converted") FROM "{CAMPAIGN_TABLE}" GROUP BY 1')
    # real queries still pass: quoted aliases, aliases defined without quotes, CTE names, string literals with quotes in them
    assert _run(table_session, f'SELECT "Channel" AS "the channel", SUM("Converted") AS s FROM "{CAMPAIGN_TABLE}" GROUP BY "the channel" ORDER BY "s" DESC').rows[0] == ("Email", 7)
    assert _run(table_session, f'WITH "t" AS (SELECT * FROM "{CAMPAIGN_TABLE}") SELECT COUNT(*) FROM "t" WHERE "Channel" = \'"Revnue" it\'\'s\'').rows == [(0,)]


# ---- real-world sheet shapes ---------------------------------------------------------------

def test_a_title_row_and_blank_row_above_the_header_keep_real_sheet_row_numbers(session_id, upload):
    from app.analytics.tablestore import session_tables
    from app.db import SessionLocal
    upload(session_id, CAMPAIGN_FILE, make_campaign_xlsx(notes=True), XLSX_TYPE)
    with SessionLocal() as db:
        [table] = session_tables(db, session_id)
        assert (table.first_row, table.last_row, table.row_count) == (4, 33, 30)     # title row 1, blank row 2, header row 3
        assert [c["name"] for c in table.columns_json] == ["Campaign_ID", "Channel", "Impressions", "Cost", "Converted"]


def test_every_sheet_of_a_workbook_is_its_own_table_and_a_question_can_name_the_sheet(llm, session_id, upload, ask):
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    sales = wb.active
    sales.title = "Sales"
    sales.append(["Region", "Amount"])
    for region, amount in [("North", 100), ("South", 250), ("North", 50)]:
        sales.append([region, amount])
    costs = wb.create_sheet("Costs")
    costs.append(["Department", "Budget"])
    costs.append(["HR", 10])
    buf = io.BytesIO(); wb.save(buf)
    document, _ = upload(session_id, "Report 2024.xlsx", buf.getvalue(), XLSX_TYPE)
    assert [t["name"] for t in document["details"]["tables"]] == ["report_2024_sales", "report_2024_costs"]

    llm.sql, llm.sql_label = 'SELECT "Region", SUM("Amount") AS total_amount FROM "report_2024_sales" GROUP BY "Region" ORDER BY total_amount DESC', "Sales by region"
    result = ask(session_id, "What is the total amount by region in the Sales sheet?")
    assert "South | 250" in result["message"]["content"] and result["citations"][0]["sheet"] == "Sales"
    assert '"report_2024_costs"' in llm.sql_prompt and '"report_2024_sales"' in llm.sql_prompt


def test_dates_booleans_currency_and_blank_cells_load_sensibly(llm, session_id, upload, ask):
    import datetime as dt
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Order Date", "Paid", "Price", "Zip"])
    ws.append([dt.datetime(2024, 3, 5), True, "$1,200.50", "01234"])
    ws.append([dt.datetime(2024, 3, 6), False, "$99.50", "02134"])
    ws.append([dt.datetime(2024, 3, 7), None, None, "10001"])
    buf = io.BytesIO(); wb.save(buf)
    upload(session_id, "orders.xlsx", buf.getvalue(), XLSX_TYPE)
    llm.sql = 'SELECT SUM("Price") AS revenue, SUM("Paid") AS paid_orders, MIN("Order_Date") AS first_day, MIN("Zip") AS zip FROM "orders"'
    answer = ask(session_id, "What is the total price and the number of paid orders?")["message"]["content"]
    assert "1300" in answer.replace(",", "") and "2024-03-05" in answer and "01234" in answer    # zip stays text, price is a number
