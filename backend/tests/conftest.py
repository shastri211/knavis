"""Test isolation.

The app reads settings from the project's real .env at import time. These variables are set
BEFORE any ``app`` import so tests never touch real API keys, the real database, or the
network, and never write into the project's data directory.
"""
import io
import json
import os
import tempfile

import pytest

os.environ["DATA_DIR"] = tempfile.mkdtemp(prefix="mrag_tests_")
for _key in ("NVIDIA_API_KEY", "GROQ_API_KEY", "OPENROUTER_API_KEY", "ASSEMBLYAI_API_KEY",
             "QDRANT_URL", "QDRANT_API_KEY", "OPENROUTER_MODEL"):
    os.environ[_key] = ""
for _limit in ("RATE_LIMIT_CHAT_PER_MINUTE", "RATE_LIMIT_UPLOAD_PER_MINUTE", "RATE_LIMIT_AUTH_PER_MINUTE"):
    os.environ[_limit] = "0"          # limits are tested on their own; the golden tests send many requests quickly
os.environ["DEFAULT_PROVIDER"] = "groq"
os.environ["DEFAULT_MODEL"] = "openai/gpt-oss-20b"

MODEL = "openai/gpt-oss-20b"
FACT = "Company data must be retained for 90 days after the contract ends."
OTHER_FACTS = "Backups are kept for 30 days. Audit logs are retained for 365 days. Deletion requests are answered within 14 days."


class FakeLLM:
    """Stands in for the provider and records every call by kind (router/answer/conversation)."""

    def __init__(self):
        self.calls: list[str] = []
        self.answer = f"{FACT} [EVIDENCE 1]"
        self.router = {"intent": "RAG_QUERY", "route": "rag"}
        self.answer_prompt = ""
        self.sql = None               # what the SQL writer returns for an analytical question
        self.sql_label = "Result"
        self.sql_raw = None           # a raw reply that replaces the JSON one (for malformed output)
        self.sql_prompt = ""

    async def chat(self, provider, model, messages, **kwargs):
        from app.providers import Response
        system = messages[0]["content"]
        if "semantic router" in system:
            self.calls.append("router")
            text = json.dumps({**self.router, "language": "English", "confidence": 0.9, "reason": "test"})
        elif "SQL query-writing component" in system:
            self.calls.append("sql")
            self.sql_prompt = messages[-1]["content"]
            text = self.sql_raw if self.sql_raw is not None else json.dumps({"sql": self.sql, "label": self.sql_label})
        elif "answer-generation component" in system:
            self.calls.append("answer")
            self.answer_prompt = messages[-1]["content"]
            text = self.answer
        else:
            self.calls.append("conversation")
            text = "Sure, happy to chat."
        return Response(text, provider, model, {"prompt_tokens": 5, "completion_tokens": 2, "total_tokens": 7})


@pytest.fixture(scope="session", autouse=True)
def database():
    """Create the schema once, so tests that touch the database do not depend on the app having started."""
    from app.db import init_db
    init_db()


@pytest.fixture
def llm(monkeypatch):
    fake = FakeLLM()
    monkeypatch.setattr("app.provider_service.raw_chat", fake.chat)
    return fake


TEST_EMAIL, TEST_PASSWORD = "tester@example.com", "correct horse battery"


def sign_in(test_client, email=TEST_EMAIL, password=TEST_PASSWORD):
    """Register (or sign in) and make every later request of this client carry the bearer token."""
    response = test_client.post("/api/auth/register", json={"email": email, "password": password})
    if response.status_code == 409:
        response = test_client.post("/api/auth/login", json={"email": email, "password": password})
    assert response.status_code in (200, 201), response.text
    test_client.headers["Authorization"] = f"Bearer {response.json()['token']}"
    return response.json()


@pytest.fixture(scope="session")
def client():
    from fastapi.testclient import TestClient
    from app.main import app
    with TestClient(app) as test_client:
        sign_in(test_client)
        yield test_client


@pytest.fixture
def other_client():
    """A second signed-in user, for ownership tests."""
    from fastapi.testclient import TestClient
    from app.main import app
    with TestClient(app) as test_client:
        sign_in(test_client, "someone.else@example.com", "another long passphrase")
        yield test_client


@pytest.fixture
def session_id(client):
    return client.post("/api/sessions", json={"title": "test"}).json()["id"]


# ---- sample documents -------------------------------------------------------------------

def make_txt(text=None):
    return (text or f"Retention Policy\n\n{FACT} {OTHER_FACTS}").encode()

def make_csv():
    return f"rule,detail\nretention,{FACT}\nbackups,Backups are kept for 30 days.\n".encode()

def make_json():
    return json.dumps({"policy": {"retention": FACT, "backups": "Backups are kept for 30 days."}}).encode()

def make_docx():
    from docx import Document
    doc = Document()
    doc.add_paragraph("Retention Policy")
    doc.add_paragraph(FACT)
    doc.add_paragraph(OTHER_FACTS)
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()

def make_pptx():
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Retention Policy"
    slide.placeholders[1].text = FACT
    buf = io.BytesIO(); prs.save(buf)
    return buf.getvalue()

def make_xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Policy"
    ws.append(["Rule", "Detail"])
    ws.append(["Retention", FACT])
    ws.append(["Backups", "Backups are kept for 30 days."])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()

def make_pdf():
    import fitz
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_textbox(fitz.Rect(72, 72, 520, 300), f"Retention Policy\n\n{FACT} {OTHER_FACTS}", fontsize=11)
    data = pdf.tobytes()
    pdf.close()
    return data

def make_html():
    return f"<html><head><title>Policy</title></head><body><h1>Retention Policy</h1><p>{FACT}</p><p>{OTHER_FACTS}</p></body></html>".encode()

def make_tsv():
    return f"rule\tdetail\nretention\t{FACT}\nbackups\tBackups are kept for 30 days.\n".encode()

def make_xml():
    return f"<policy><retention>{FACT}</retention><backups>Backups are kept for 30 days.</backups></policy>".encode()

def make_srt():
    return f"1\n00:00:01,000 --> 00:00:06,000\n{FACT}\n\n2\n00:00:07,000 --> 00:00:09,000\nBackups are kept for 30 days.\n".encode()

def make_eml():
    from email.message import EmailMessage
    msg = EmailMessage()
    msg["From"], msg["To"], msg["Subject"] = "a@example.com", "b@example.com", "Retention policy"
    msg.set_content(f"{FACT} {OTHER_FACTS}")
    return bytes(msg)

def make_yaml():
    return f"policy:\n  retention: {FACT}\n  backups: Backups are kept for 30 days.\n".encode()

CAMPAIGN_ROWS = (
    # (channel, converted) per campaign row: Email converts 7 of 12, Search 5 of 8, Social 3 of 10
    [("Email", 1)] * 7 + [("Email", 0)] * 5 + [("Search", 1)] * 5 + [("Search", 0)] * 3 + [("Social", 1)] * 3 + [("Social", 0)] * 7
)

def make_campaign_xlsx(title="Week1", notes=False):
    """A marketing sheet in the style of Campaign_Data_Week1_new.xlsx: Converted is a 1/0 flag per row."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    if notes:
        ws.append(["Weekly campaign export"])
        ws.append([])
    ws.append(["Campaign_ID", "Channel", "Impressions", "Cost", "Converted"])
    for i, (channel, converted) in enumerate(CAMPAIGN_ROWS, 1):
        ws.append([f"C{i:03d}", channel, 1000 + i * 10, round(10.5 + i, 2), converted])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()

CAMPAIGN_FILE = "Campaign_Data_Week1_new.xlsx"
CAMPAIGN_TABLE = "campaign_data_week1_new"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SAMPLES = {
    "txt": ("policy.txt", make_txt, "text/plain"),
    "md": ("policy.md", make_txt, "text/markdown"),
    "csv": ("policy.csv", make_csv, "text/csv"),
    "json": ("policy.json", make_json, "application/json"),
    "docx": ("policy.docx", make_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "pptx": ("policy.pptx", make_pptx, "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    "xlsx": ("policy.xlsx", make_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "pdf": ("policy.pdf", make_pdf, "application/pdf"),
    "html": ("policy.html", make_html, "text/html"),
    "tsv": ("policy.tsv", make_tsv, "text/tab-separated-values"),
    "xml": ("policy.xml", make_xml, "application/xml"),
    "srt": ("policy.srt", make_srt, "application/x-subrip"),
    "eml": ("policy.eml", make_eml, "message/rfc822"),
    "yaml": ("policy.yaml", make_yaml, "application/yaml"),
}


def table_files(session_id):
    """The spreadsheet table files of a chat that exist on this host (the per-document files, not derived views)."""
    from app.analytics.tablestore import session_prefix
    from app.config import settings
    return sorted(p for p in (settings.data_dir / session_prefix(session_id)).glob("*.sqlite") if not p.name.startswith("view-"))


@pytest.fixture
def upload(client):
    """Upload bytes to a session and return (document, job). Background ingestion has finished on return."""
    def _upload(session_id, name, data, content_type="application/octet-stream"):
        response = client.post("/api/uploads", data={"session_id": session_id}, files={"file": (name, data, content_type)})
        assert response.status_code == 200, response.text
        body = response.json()
        job = client.get(f"/api/jobs/{body['job']['id']}").json()
        document = next(d for d in client.get(f"/api/sessions/{session_id}/documents").json() if d["id"] == body["document"]["id"])
        return document, job
    return _upload


@pytest.fixture
def ask(client):
    def _ask(session_id, question, **extra):
        response = client.post("/api/chat", json={"session_id": session_id, "content": question, "provider": "groq", "model": MODEL, **extra})
        assert response.status_code == 200, response.text
        return response.json()
    return _ask
