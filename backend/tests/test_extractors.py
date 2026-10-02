import io
import shutil

import fitz
import pytest

from app.ingest.elements import (
    HEADING, PARAGRAPH, RECORD, SLIDE, SUMMARY, TABLE, TRANSCRIPT, ExtractionError,
)
from app.ingest.extract import extract_document
from app.ingest.registry import ALLOWED_EXTENSIONS


def kinds(result):
    return [e.kind for e in result.elements]


# ---- PDF ---------------------------------------------------------------------------------

def make_pdf(path, scanned_page=False):
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Retention Policy", fontsize=22)
    page.insert_textbox(fitz.Rect(72, 100, 520, 200), "Company data must be retained for 90 days after the contract ends. " * 3, fontsize=11)
    for r, row in enumerate([["Item", "Days", "Owner"], ["Backups", "30", "IT"], ["Logs", "365", "Sec"]]):
        for c, value in enumerate(row):
            cell = fitz.Rect(72 + c * 120, 260 + r * 24, 72 + (c + 1) * 120, 284 + r * 24)
            page.draw_rect(cell)
            page.insert_text((cell.x0 + 4, cell.y0 + 16), value, fontsize=10)
    pdf.new_page()                                   # a truly blank page
    if scanned_page:
        scan = pdf.new_page()
        pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 400, 500), False)
        pixmap.clear_with(200)
        scan.insert_image(scan.rect, pixmap=pixmap)  # page-sized image, no text layer
    pdf.save(path)
    pdf.close()


def test_pdf_headings_paragraphs_and_tables(tmp_path):
    make_pdf(tmp_path / "a.pdf")
    result = extract_document(tmp_path / "a.pdf")
    heading = result.elements[0]
    assert (heading.kind, heading.text, heading.level, heading.page) == (HEADING, "Retention Policy", 1, 1)
    table = next(e for e in result.elements if e.kind == TABLE)
    assert table.rows == [["Item", "Days", "Owner"], ["Backups", "30", "IT"], ["Logs", "365", "Sec"]]
    # the table's text appears once, as a table, not again as loose paragraphs
    assert sum("Backups" in e.text for e in result.elements) == 1
    assert result.info["tables"] == 1 and result.info["pages"] == 2 and result.info["ocr_pages"] == []


def test_scanned_page_is_queued_for_ocr_but_a_blank_page_is_not(tmp_path):
    make_pdf(tmp_path / "s.pdf", scanned_page=True)
    result = extract_document(tmp_path / "s.pdf")
    assert result.info["ocr_pages"] == [3]            # page 2 is blank: no OCR call is wasted on it
    assert result.info["estimated_ocr_calls"] == 1
    assert result.info["figure_pages"] == [{"page": 3, "count": 1}]


def test_password_protected_and_corrupt_pdfs_give_clear_errors(tmp_path):
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "secret")
    pdf.save(tmp_path / "locked.pdf", encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw")
    with pytest.raises(ExtractionError, match="password"):
        extract_document(tmp_path / "locked.pdf")
    (tmp_path / "bad.pdf").write_bytes(b"not a pdf at all")
    with pytest.raises(ExtractionError):
        extract_document(tmp_path / "bad.pdf")


def test_combined_pdf_is_split_into_logical_documents(tmp_path):
    pdf = fitz.open()
    for text in ("Some report text.", "Invoice number 55 for services", "Invoice continued here"):
        pdf.new_page().insert_text((72, 72), text)
    pdf.save(tmp_path / "bundle.pdf")
    result = extract_document(tmp_path / "bundle.pdf")
    assert result.info["logical_documents"] == 2
    assert {e.group for e in result.elements} == {"logical:1", "logical:2"}


# ---- Office ------------------------------------------------------------------------------

def test_docx_keeps_document_order_headings_and_tables(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_heading("Policy", 1)
    doc.add_paragraph("Intro paragraph.")
    table = doc.add_table(rows=2, cols=2)
    for r, row in enumerate([["Item", "Days"], ["Backups", "30"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    doc.add_paragraph("Closing paragraph.")
    doc.add_paragraph("a bullet", style="List Bullet")
    doc.save(tmp_path / "a.docx")
    result = extract_document(tmp_path / "a.docx")
    assert kinds(result) == [HEADING, PARAGRAPH, TABLE, PARAGRAPH, PARAGRAPH]
    assert result.elements[0].level == 1 and result.elements[2].rows == [["Item", "Days"], ["Backups", "30"]]
    assert result.elements[4].text == "- a bullet"


def test_pptx_slides_include_title_body_notes_tables_and_picture_counts(tmp_path):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches
    Image.new("RGB", (40, 40), "red").save(tmp_path / "pic.png")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Retention"
    slide.placeholders[1].text = "Keep data 90 days"
    slide.notes_slide.notes_text_frame.text = "Mention the exceptions"
    grid = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
    for r, row in enumerate([["Item", "Days"], ["Logs", "365"]]):
        for c, value in enumerate(row):
            grid.cell(r, c).text = value
    only_picture = prs.slides.add_slide(prs.slide_layouts[6])
    only_picture.shapes.add_picture(str(tmp_path / "pic.png"), Inches(1), Inches(1))
    prs.save(tmp_path / "a.pptx")

    result = extract_document(tmp_path / "a.pptx")
    slide_el = next(e for e in result.elements if e.kind == SLIDE)
    assert slide_el.slide == 1 and "Retention" in slide_el.text and "Keep data 90 days" in slide_el.text
    assert "Speaker notes: Mention the exceptions" in slide_el.text
    assert next(e for e in result.elements if e.kind == TABLE).rows == [["Item", "Days"], ["Logs", "365"]]
    assert result.info["picture_only_slides"] == [2] and result.info["figure_slides"] == [{"slide": 2, "count": 1}]


# ---- spreadsheets and delimited text -----------------------------------------------------

def test_xlsx_sheet_summary_and_row_groups_repeat_the_header(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Amount", "Date"])
    import datetime
    for i in range(600):
        ws.append([f"region-{i % 7}", i * 1.5, datetime.date(2024, 1, 1) + datetime.timedelta(days=i % 28)])
    wb.create_sheet("Empty")
    wb.save(tmp_path / "a.xlsx")

    result = extract_document(tmp_path / "a.xlsx")
    summary = next(e for e in result.elements if e.kind == SUMMARY)
    assert "600 data rows" in summary.text and "Region (text)" in summary.text
    assert "Amount (number)" in summary.text and "Date (date)" in summary.text
    groups = [e for e in result.elements if e.kind == TABLE]
    assert len(groups) > 3
    assert all(g.rows[0] == ["Region", "Amount", "Date"] and g.sheet == "Sales" for g in groups)
    assert all(len(g.text) <= 1300 for g in groups)                      # no giant single chunk
    assert groups[0].locator.startswith("sheet Sales, rows 2-")
    assert sum(len(g.rows) - 1 for g in groups) == 600                    # every row kept exactly once


def test_csv_dialect_and_encoding_detection(tmp_path):
    (tmp_path / "semi.csv").write_bytes("nom;prix\ncafé;3,5\nthé;4\n".encode("cp1252"))
    result = extract_document(tmp_path / "semi.csv")
    table = next(e for e in result.elements if e.kind == TABLE)
    assert table.rows == [["nom", "prix"], ["café", "3,5"], ["thé", "4"]]
    (tmp_path / "t.tsv").write_text("a\tb\n1\t2\n", encoding="utf-8")
    assert next(e for e in extract_document(tmp_path / "t.tsv").elements if e.kind == TABLE).rows == [["a", "b"], ["1", "2"]]


# ---- text family -------------------------------------------------------------------------

def test_json_is_flattened_so_every_line_stands_alone(tmp_path):
    (tmp_path / "a.json").write_text('{"policy": {"retention": "90 days", "tags": ["a", "b"]}, "items": [{"id": 1}, {"id": 2}]}', encoding="utf-8")
    (record,) = extract_document(tmp_path / "a.json").elements
    assert record.kind == RECORD
    assert record.text.splitlines() == ["policy.retention: 90 days", "policy.tags: a, b", "items[0].id: 1", "items[1].id: 2"]
    (tmp_path / "bad.json").write_text('{"a": ', encoding="utf-8")
    with pytest.raises(ExtractionError, match="line"):
        extract_document(tmp_path / "bad.json")


def test_markdown_headings_and_code_fences(tmp_path):
    (tmp_path / "a.md").write_text("# Title\n\nIntro text.\n\n## Part\n\n```\n# not a heading\ncode\n```\n", encoding="utf-8")
    result = extract_document(tmp_path / "a.md")
    assert [(e.kind, e.text, e.level) for e in result.elements][:3] == [(HEADING, "Title", 1), (PARAGRAPH, "Intro text.", None), (HEADING, "Part", 2)]
    assert "# not a heading" in result.elements[-1].text and result.elements[-1].kind == PARAGRAPH


def test_html_drops_scripts_and_keeps_structure(tmp_path):
    (tmp_path / "a.html").write_text(
        "<html><head><title>T</title><style>p{}</style></head><body><script>evil()</script>"
        "<h1>Heading</h1><p>First para.</p><table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>"
        "<p>Last para.</p></body></html>", encoding="utf-8")
    result = extract_document(tmp_path / "a.html")
    assert kinds(result) == [HEADING, PARAGRAPH, TABLE, PARAGRAPH]
    assert "evil" not in " ".join(e.text for e in result.elements) and result.info["title"] == "T"
    assert result.elements[2].rows == [["A", "B"], ["1", "2"]]


def test_xml_is_flattened_with_attributes(tmp_path):
    (tmp_path / "a.xml").write_text('<policy id="7"><retention unit="days">90</retention></policy>', encoding="utf-8")
    (record,) = extract_document(tmp_path / "a.xml").elements
    assert record.text.splitlines() == ["policy@id: 7", "policy.retention@unit: days", "policy.retention: 90"]


def test_subtitles_keep_timestamps(tmp_path):
    (tmp_path / "a.srt").write_text("1\n00:00:01,000 --> 00:00:03,500\nHello <i>there</i>\n\n2\n00:01:05,000 --> 00:01:08,000\nSecond cue\n", encoding="utf-8")
    cues = extract_document(tmp_path / "a.srt").elements
    assert [c.kind for c in cues] == [TRANSCRIPT, TRANSCRIPT]
    assert cues[0].text == "Hello there" and cues[0].meta == {"start_s": 1.0, "end_s": 3.5} and cues[0].locator == "00:01-00:03"
    (tmp_path / "a.vtt").write_text("WEBVTT\n\n00:05.000 --> 00:07.000\nShort form\n", encoding="utf-8")
    assert extract_document(tmp_path / "a.vtt").elements[0].meta["start_s"] == 5.0


def test_email_headers_body_and_attachment_names(tmp_path):
    from email.message import EmailMessage
    msg = EmailMessage()
    msg["From"], msg["To"], msg["Subject"], msg["Date"] = "a@x.com", "b@x.com", "Policy update", "Mon, 1 Jan 2024 10:00:00 +0000"
    msg.set_content("Retention is now 90 days.")
    msg.add_attachment(b"%PDF", maintype="application", subtype="pdf", filename="policy.pdf")
    (tmp_path / "a.eml").write_bytes(bytes(msg))
    result = extract_document(tmp_path / "a.eml")
    text = "\n".join(e.text for e in result.elements)
    assert "Subject: Policy update" in text and "Retention is now 90 days." in text
    assert "Attachments (not indexed): policy.pdf" in text and result.info["attachments"] == ["policy.pdf"]


# ---- registry, images/audio, legacy ------------------------------------------------------

def test_registry_lists_the_supported_formats():
    for ext in (".pdf", ".docx", ".pptx", ".xlsx", ".xlsm", ".csv", ".tsv", ".md", ".txt", ".json", ".html", ".xml", ".srt", ".vtt",
                ".eml", ".png", ".jpg", ".mp3", ".wav", ".doc", ".ppt", ".xls", ".yaml", ".py"):
        assert ext in ALLOWED_EXTENSIONS, ext
    assert ".exe" not in ALLOWED_EXTENSIONS and ".zip" not in ALLOWED_EXTENSIONS


def test_images_and_audio_report_the_specialist_they_need_without_calling_it(tmp_path):
    (tmp_path / "a.png").write_bytes(b"\x89PNG\r\n\x1a\n")
    (tmp_path / "a.mp3").write_bytes(b"ID3")
    image, audio = extract_document(tmp_path / "a.png"), extract_document(tmp_path / "a.mp3")
    assert image.elements == [] and image.info["estimated_ocr_calls"] == 1
    assert audio.elements == [] and audio.info["estimated_asr_calls"] == 1


def test_legacy_formats_explain_what_is_missing_when_libreoffice_is_absent(tmp_path, monkeypatch):
    monkeypatch.setattr("app.ingest.extractors.legacy.find_soffice", lambda: None)
    (tmp_path / "old.doc").write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(ExtractionError, match="LibreOffice"):
        extract_document(tmp_path / "old.doc")


def test_legacy_formats_are_converted_then_extracted(tmp_path, monkeypatch):
    from docx import Document
    modern = tmp_path / "made.docx"
    doc = Document(); doc.add_paragraph("Converted content"); doc.save(modern)

    def fake_convert(path, target_ext, out_dir):
        out_dir.mkdir(parents=True, exist_ok=True)
        converted = out_dir / (path.stem + target_ext)
        shutil.copy(modern, converted)
        return converted

    monkeypatch.setattr("app.ingest.extract.convert_legacy", fake_convert)
    (tmp_path / "old.doc").write_bytes(b"\xd0\xcf\x11\xe0")
    result = extract_document(tmp_path / "old.doc", tmp_path / "work")
    assert result.elements[0].text == "Converted content" and result.info["converted_from"] == ".doc"


def test_title_rows_above_the_header_and_blank_rows_keep_real_row_numbers(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(["Quarterly report"])             # row 1: title, not a header
    ws.append([])                               # row 2: blank
    ws.append(["Region", "Amount"])             # row 3: the real header
    ws.append(["north", 10])                    # row 4
    ws.append([])                               # row 5: blank
    ws.append(["south", 20])                    # row 6
    wb.save(tmp_path / "a.xlsx")

    result = extract_document(tmp_path / "a.xlsx")
    preamble = next(e for e in result.elements if e.kind == PARAGRAPH)
    assert preamble.text == "Quarterly report" and preamble.locator == "sheet Report, row 1"
    summary = next(e for e in result.elements if e.kind == SUMMARY)
    assert "Region (text), Amount (number)" in summary.text and summary.meta["header_row"] == 3
    (group,) = [e for e in result.elements if e.kind == TABLE]
    assert group.rows == [["Region", "Amount"], ["north", "10"], ["south", "20"]]
    assert group.locator == "sheet Report, rows 4-6"        # where a person would look in the sheet
